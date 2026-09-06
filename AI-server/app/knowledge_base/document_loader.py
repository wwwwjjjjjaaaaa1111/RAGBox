"""根据知识库文件类型选择合适的文档加载器。"""

from __future__ import annotations

from pathlib import Path

from langchain_community.document_loaders import BSHTMLLoader, Docx2txtLoader, PyPDFLoader, TextLoader
from langchain_core.documents import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".html", ".htm", ".xlsx"}

SUPPORTED_EXTENSIONS_LABEL = "PDF / DOCX / TXT / MD / HTML / XLSX"


def _load_xlsx(path: Path) -> list[Document]:
    """用 openpyxl 加载 xlsx：每个 sheet 生成一个 Document，行内单元格用制表符连接。"""

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    documents: list[Document] = []

    try:
        for sheet in workbook.worksheets:
            lines: list[str] = [f"# {sheet.title}"]
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if cell is None else str(cell).strip() for cell in row]
                # 跳过整行空白。
                if not any(cells):
                    continue
                lines.append("\t".join(cells))

            content = "\n".join(lines).strip()
            if content:
                documents.append(
                    Document(
                        page_content=content,
                        metadata={"source": str(path), "sheet": sheet.title},
                    )
                )
    finally:
        workbook.close()

    return documents


def load_knowledge_documents(file_path: str) -> list[Document]:
    """将受支持的知识库文件加载为 LangChain 文档对象。"""

    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix} (supported: {SUPPORTED_EXTENSIONS_LABEL})")

    # PDF 按页拆成多个 Document，其余格式通常作为单个逻辑文档处理。
    if suffix == ".pdf":
        return PyPDFLoader(str(path), mode="page").load()
    if suffix == ".docx":
        return Docx2txtLoader(str(path)).load()
    if suffix in (".html", ".htm"):
        return BSHTMLLoader(str(path)).load()
    if suffix == ".xlsx":
        return _load_xlsx(path)
    return TextLoader(str(path), encoding="utf-8").load()
